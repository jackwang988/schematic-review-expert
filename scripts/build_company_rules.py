# -*- coding: utf-8 -*-
"""
build_company_rules.py
将《公司HW checklist_汇总.xlsx》转换为结构化公司规范库。

输出：
  knowledge_base/company_rules/<module>.md      每个功能模块一个规则文件
  knowledge_base/company_rules/index.json        全部规则条目索引（供 kb_search 检索）
  knowledge_base/review_cases/classic_cases.md    经典案例经验库

用法：
  python build_company_rules.py "<checklist.xlsx>" "<skill_root>"
若不传参，默认使用脚本相对路径推断的 skill 根目录，并从 uploads/ 下自动查找 checklist。
"""
import sys, os, json, re

def load_wb(path):
    import openpyxl
    return openpyxl.load_workbook(path, data_only=True)

# 严重度中文 -> 报告标识映射
SEV_MAP = {"高": "严重", "中": "一般", "低": "提示", "": "一般"}

# SCH 检查大类 -> 功能模块归类（用于按需加载）
def classify_module(check_category, content):
    c = (check_category or "") + (content or "")
    if "电源" in check_category: return "电源设计"
    if "射频" in check_category: return "射频电路"
    if "对外接口" in check_category or "接口" in check_category: return "接口防护"
    if "芯片选型" in check_category: return "芯片选型"
    if "注释" in check_category: return "原理图注释"
    if "芯片外围" in check_category: return "芯片外围电路"
    return "通用"

def parse_overall(ws):
    """解析『硬件设计总体原则』sheet。列结构：B序号 C SCH/PCB D检查项目 E检查内容 F严重度"""
    rules = []
    cur_scb = ""       # SCH / PCB
    cur_category = ""   # 检查项目大类
    for row in ws.iter_rows(min_row=4, values_only=True):
        # 列: A(0)空 B(1)序号 C(2)SCH/PCB D(3)检查项目 E(4)检查内容 F(5)严重度
        idx = row[1]
        scb = (str(row[2]).strip() if row[2] else "")
        category = (str(row[3]).strip() if row[3] else "")
        content = (str(row[4]).strip() if row[4] else "")
        sev = (str(row[5]).strip() if row[5] else "")
        if scb: cur_scb = scb
        if category: cur_category = category
        if not content: continue
        if idx is None: continue
        rules.append({
            "no": str(idx).strip(),
            "domain": cur_scb or "SCH",           # SCH / PCB
            "category": cur_category,              # 检查大类
            "module": classify_module(cur_category, content),
            "content": re.sub(r'\s+\n', '\n', content),
            "severity": SEV_MAP.get(sev, "一般"),
        })
    return rules

def parse_module_sheet(ws, module_name):
    """解析摄像头/SE重点模块 sheet，抽取通用设计规则列与历史经验。"""
    rules = []
    n = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        cells = [(str(c).strip() if c is not None else "") for c in row]
        text_cells = [c for c in cells if c]
        if not text_cells:
            continue
        # 合并该行非空文本作为一条规则（模块表结构不规整，采用整行文本策略）
        content = " ｜ ".join(text_cells)
        # 过滤纯表头/分类标记行
        if content in ("原理设计 ｜ PCB layout ｜ 其他 ｜ 项目归属 ｜ 历史问题",):
            continue
        if content.startswith("模块 ｜ 设计分类"):
            continue
        n += 1
        rules.append({
            "no": f"{module_name[:2]}-{n}",
            "domain": "SCH/PCB",
            "category": module_name,
            "module": module_name,
            "content": content,
            "severity": "一般",
        })
    return rules

def parse_cases(ws):
    lines = []
    for row in ws.iter_rows(values_only=True):
        cells = [(str(c).strip() if c is not None else "") for c in row]
        text = " ".join([c for c in cells if c])
        if text.strip():
            lines.append(text.strip())
    return lines

def write_module_md(out_dir, module, rules):
    safe = module.replace("/", "_").replace(" ", "")
    path = os.path.join(out_dir, f"{safe}.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"# 公司规范库 · 模块：{module}\n\n")
        f.write("> 来源：用户导入的硬件设计 checklist（版本请按原表核实）\n\n")
        f.write("| 条款号 | 域 | 检查大类 | 默认严重度 | 约束内容 / 判定标准 |\n")
        f.write("|--------|----|----------|-----------|--------------------|\n")
        for r in rules:
            content = r["content"].replace("\n", "<br>").replace("|", "/")
            f.write(f"| {r['no']} | {r['domain']} | {r['category']} | {r['severity']} | {content} |\n")
    return path

def main():
    args = sys.argv[1:]
    script_dir = os.path.dirname(os.path.abspath(__file__))
    skill_root = os.path.dirname(script_dir)
    xlsx = None
    if len(args) >= 1: xlsx = args[0]
    if len(args) >= 2: skill_root = args[1]
    if not xlsx:
        # 从 uploads/ 找 checklist
        up = os.path.join(skill_root, "uploads")
        for fn in os.listdir(up) if os.path.isdir(up) else []:
            if fn.lower().endswith(".xlsx") and ("checklist" in fn.lower() or "checklist" in fn):
                xlsx = os.path.join(up, fn); break
    if not xlsx or not os.path.exists(xlsx):
        print("ERROR: 未找到 checklist xlsx，请传入路径。"); sys.exit(1)

    wb = load_wb(xlsx)
    company_dir = os.path.join(skill_root, "knowledge_base", "company_rules")
    cases_dir = os.path.join(skill_root, "knowledge_base", "review_cases")
    os.makedirs(company_dir, exist_ok=True)
    os.makedirs(cases_dir, exist_ok=True)

    all_rules = []
    # 1. 总体原则
    for sn in wb.sheetnames:
        if "总体原则" in sn:
            all_rules += parse_overall(wb[sn])
    # 2. 重点模块
    for sn in wb.sheetnames:
        if "摄像头" in sn:
            all_rules += parse_module_sheet(wb[sn], "摄像头模块")
        elif "SE" in sn:
            all_rules += parse_module_sheet(wb[sn], "SE安全模块")

    # 按模块分组写 md
    by_module = {}
    for r in all_rules:
        by_module.setdefault(r["module"], []).append(r)
    written = []
    for module, rules in by_module.items():
        written.append(write_module_md(company_dir, module, rules))

    # index.json
    index = {
        "source": os.path.basename(xlsx),
        "version": "V0.1",
        "total": len(all_rules),
        "modules": sorted(by_module.keys()),
        "rules": all_rules,
    }
    idx_path = os.path.join(company_dir, "index.json")
    with open(idx_path, "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=2)

    # 经典案例
    for sn in wb.sheetnames:
        if "案例" in sn:
            cases = parse_cases(wb[sn])
            cpath = os.path.join(cases_dir, "classic_cases.md")
            with open(cpath, "w", encoding="utf-8") as f:
                f.write("# 评审经验库 · 经典案例\n\n")
                f.write("> 来源：用户导入 checklist 中的案例\n\n")
                for line in cases:
                    f.write(f"- {line}\n")

    print(f"OK 总规则数={len(all_rules)}  模块数={len(by_module)}")
    print("模块:", ", ".join(f"{m}({len(r)})" for m,r in by_module.items()))
    print("index:", idx_path)

if __name__ == "__main__":
    main()
